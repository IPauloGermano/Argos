from __future__ import annotations
import os
import re
import unicodedata
from datetime import datetime
from typing import Optional

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".rtf", ".md"}
ALLOWED_MIMETYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "text/plain",
    "text/markdown",
    "text/x-markdown",
    "application/rtf",
    "text/rtf",
    "text/x-python",
    "application/octet-stream"  # alguns clientes enviam binários como octet-stream
}

SKILLS_CANONICAL = {
    # Backend & Linguagens
    "python": "Python", "django": "Django", "django orm": "Django ORM", "fastapi": "FastAPI",
    "flask": "Flask", "sql": "SQL", "sqlite": "SQLite", "postgresql": "PostgreSQL",
    "postgres": "PostgreSQL", "mysql": "MySQL", "mongodb": "MongoDB", "redis": "Redis",
    "elasticsearch": "Elasticsearch", "java": "Java", "spring boot": "Spring Boot", "spring": "Spring Boot",
    "c#": "C#", ".net": ".NET", "asp.net": "ASP.NET", "node.js": "Node.js", "nodejs": "Node.js",
    "node": "Node.js", "express": "Express", "nest.js": "NestJS", "nestjs": "NestJS",
    "php": "PHP", "laravel": "Laravel", "golang": "Go", "go": "Go", "rust": "Rust",
    "ruby": "Ruby", "ruby on rails": "Ruby on Rails", "rails": "Ruby on Rails",
    "c++": "C++", "c": "C", "kotlin": "Kotlin", "swift": "Swift",
    "typescript": "TypeScript", "javascript": "JavaScript",

    # Frontend & Mobile
    "react": "React", "react.js": "React", "reactjs": "React", "next.js": "Next.js",
    "nextjs": "Next.js", "vue": "Vue.js", "vue.js": "Vue.js", "vuejs": "Vue.js",
    "angular": "Angular", "svelte": "Svelte", "html5": "HTML5", "html": "HTML5",
    "css3": "CSS3", "css": "CSS3", "tailwind": "Tailwind CSS", "tailwind css": "Tailwind CSS",
    "tailwindcss": "Tailwind CSS", "bootstrap": "Bootstrap", "sass": "Sass",
    "redux": "Redux", "flutter": "Flutter", "react native": "React Native",

    # Arquitetura, APIs & Dados
    "apis rest": "APIs REST", "api rest": "APIs REST", "rest": "APIs REST",
    "restful": "APIs REST", "graphql": "GraphQL", "grpc": "gRPC", "microservices": "Microsserviços",
    "microsserviços": "Microsserviços", "modelagem relacional": "Modelagem Relacional",
    "modelagem de dados": "Modelagem Relacional", "bancos relacionais": "Modelagem Relacional",
    "sqlalchemy": "SQLAlchemy", "prisma": "Prisma", "hibernate": "Hibernate",
    "pandas": "Pandas", "numpy": "NumPy", "scikit-learn": "Scikit-Learn", "power bi": "Power BI",

    # DevOps, Cloud & Infra
    "git": "Git", "github": "GitHub", "gitlab": "GitLab", "linux": "Linux",
    "docker": "Docker", "docker compose": "Docker Compose", "kubernetes": "Kubernetes",
    "k8s": "Kubernetes", "helm": "Helm", "terraform": "Terraform", "ansible": "Ansible",
    "ci/cd": "CI/CD", "github actions": "GitHub Actions", "jenkins": "Jenkins",
    "aws": "AWS", "azure": "Azure", "gcp": "GCP", "google cloud": "GCP", "nginx": "Nginx",

    # Testes & Qualidade
    "testes unitários": "Testes Unitários", "unit testing": "Testes Unitários", "unit tests": "Testes Unitários",
    "testes de integração": "Testes de Integração", "integration tests": "Testes de Integração",
    "tdd": "TDD", "pytest": "Pytest", "unittest": "Unittest", "jest": "Jest", "cypress": "Cypress",

    # Metodologias & IA
    "prompt engineering": "Prompt Engineering", "ia generativa": "IA Generativa",
    "inteligência artificial": "Inteligência Artificial", "artificial intelligence": "Inteligência Artificial",
    "machine learning": "Machine Learning", "deep learning": "Deep Learning",
    "llms": "LLMs", "llm": "LLMs", "langchain": "LangChain",
    "scrum": "Scrum", "kanban": "Kanban", "metodologias ágeis": "Metodologias Ágeis",
    "agile": "Metodologias Ágeis", "clean code": "Clean Code", "clean architecture": "Clean Architecture",
    "arquitetura limpa": "Clean Architecture", "solid": "SOLID", "12-factor": "12-Factor App",
    "12-factor app": "12-Factor App"
}

SECTION_PATTERNS = {
    "summary": [
        r"\b(resumo\s+profissional|resumo|sobre\s+mim|perfil\s+profissional|perfil|apresenta[çc][ãa]o|objetivo\s+profissional|objetivo|professional\s+summary|summary|about\s+me|profile|career\s+objective|executive\s+summary|resumen\s+profesional|resumen)\b"
    ],
    "experience": [
        r"\b(experi[êe]ncia\s+profissional|experi[êe]ncias?\s+profissionais|experi[êe]ncia|experi[êe]ncias|hist[óo]rico\s+profissional|atua[çc][ãa]o\s+profissional|trajet[óo]ria\s+profissional|work\s+experience|professional\s+experience|employment\s+history|experience|experiencia\s+laboral)\b"
    ],
    "education": [
        r"\b(forma[çc][ãa]o\s+acad[êe]mica|forma[çc][ãa]o|educa[çc][ãa]o|escolaridade|ensino\s+superior|educaci[óo]n|academic\s+background|education|degrees?)\b"
    ],
    "skills": [
        r"\b(conhecimentos?\s+t[ée]cnicos?|habilidades?\s+t[ée]cnicas?|habilidades|compet[êe]ncias|conhecimentos|tecnologias|stacks?|ferramentas|technical\s+skills|core\s+competencies|skills|technologies|tools|hard\s+skills)\b"
    ],
    "languages": [
        r"\b(idiomas?|l[íi]nguas?|languages?|language\s+proficiency)\b"
    ],
    "projects": [
        r"\b(projetos?\s+relevantes?|projetos?\s+pessoais|projetos?|pr[áa]tica\s+de\s+desenvolvimento|projects?|key\s+projects?|personal\s+projects?|portfolio|portf[óo]lio)\b"
    ]
}


def validate_upload(filename: str, content_type: str, size: int, max_mb: int = 10) -> None:
    """Valida tamanho, presença de dados e formato do arquivo de currículo."""
    if size == 0:
        raise ValueError("Arquivo vazio")

    if size > max_mb * 1024 * 1024:
        raise ValueError(f"Arquivo maior que {max_mb}MB")

    ext = os.path.splitext(filename.lower())[1] if filename else ""
    if ext:
        if ext not in ALLOWED_EXTENSIONS:
            raise ValueError(f"Formato não suportado ('{ext}'). Envie arquivos PDF, DOCX, TXT, RTF ou MD")
    elif content_type not in ALLOWED_MIMETYPES:
        raise ValueError(f"Formato não suportado ('{content_type}'). Envie arquivos PDF, DOCX, TXT, RTF ou MD")


def _clean_extracted_text(text: str) -> str:
    """Normaliza ligaduras de PDF, quebras de hífens, marcadores e espaçamentos."""
    if not text:
        return ""
    # Substituição de ligaduras tipográficas comuns em PDFs
    ligatures = {
        "ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl",
        "\xad": "", "\u200b": ""
    }
    for lig, rep in ligatures.items():
        text = text.replace(lig, rep)

    # Junta palavras hifenizadas no fim da linha (ex: desenvolvi-\n mento -> desenvolvimento)
    text = re.sub(r"(\w+)-\s*\n\s*(\w+)", r"\1\2", text)

    # Normaliza marcadores gráficos para marcadores textuais padrão
    text = re.sub(r"[\u2022\u25cf\u25aa\u25b8\u25b9\u25cb\u25e6\u2714\u2713]", "•", text)

    # Remove repetição excessiva de quebras de linha
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_from_pdf(path: str) -> str:
    """Extrai texto de arquivo PDF utilizando pypdf com suporte a layout multi-colunas."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(path)
        if len(reader.pages) == 0:
            raise ValueError("O arquivo PDF não contém páginas")

        parts = []
        for page in reader.pages:
            page_text = ""
            # Tenta primeiramente o modo layout (mantém colunas intactas em currículos Canva, Overleaf, etc.)
            try:
                page_text = page.extract_text(extraction_mode="layout") or ""
            except Exception:
                page_text = ""
            if not page_text.strip():
                page_text = page.extract_text() or ""
            if page_text.strip():
                parts.append(page_text)

        text = "\n\n".join(parts).strip()
        text = _clean_extracted_text(text)
        if not text:
            raise ValueError(
                "O arquivo PDF não contém texto legível por máquina (pode ser uma imagem escaneada). "
                "Por favor, envie uma versão gerada em texto por Word, Google Docs, Canva ou TXT."
            )
        return text
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Arquivo PDF corrompido ou ilegível: {e}")


def extract_text_from_docx(path: str) -> str:
    """Extrai texto de DOCX cobrindo parágrafos, tabelas e caixas de texto (textboxes)."""
    try:
        import docx
        doc = docx.Document(path)
        parts = []

        # 1. Parágrafos principais
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())

        # 2. Tabelas (muito comuns em modelos com colunas ou cabeçalhos laterais)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)

        # 3. Caixas de texto e formas de desenho no XML do corpo do documento
        try:
            for elem in doc.element.body.iter():
                if elem.tag.endswith("txbxContent"):
                    box_texts = [node.text.strip() for node in elem.iter() if node.tag.endswith("t") and node.text and node.text.strip()]
                    if box_texts:
                        parts.append(" ".join(box_texts))
        except Exception:
            pass

        text = "\n".join(parts).strip()
        text = _clean_extracted_text(text)
        if not text:
            raise ValueError("O documento DOCX não contém texto legível")
        return text
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Arquivo DOCX corrompido ou ilegível: {e}")


def extract_text_from_rtf(path: str) -> str:
    """Extrai texto limpo de arquivo RTF (Rich Text Format)."""
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            raw = f.read()
    except Exception:
        with open(path, "r", encoding="latin-1", errors="ignore") as f:
            raw = f.read()

    # Remove grupos de cabeçalhos e tabelas de fontes/cores do RTF
    text = re.sub(r"\{\*?\\[^{}]+(?:\{[^{}]*\})*\}", "", raw)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
    text = text.replace("\\\\", "\\").replace("\\{", "{").replace("\\}", "}")
    text = re.sub(r"[{}]", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        raise ValueError("O arquivo RTF está vazio ou ilegível")
    return text


def extract_text_from_txt(path: str) -> str:
    """Extrai texto de arquivo TXT ou Markdown suportando UTF-8, Latin-1 e CP1252."""
    encodings = ["utf-8", "latin-1", "cp1252", "iso-8859-1"]
    text = ""
    for enc in encodings:
        try:
            with open(path, "r", encoding=enc) as f:
                text = f.read().strip()
            if text:
                break
        except Exception:
            continue

    if not text:
        raise ValueError("O arquivo de texto está vazio")
    return _clean_extracted_text(text)


def extract_text_from_upload(path: str, content_type: str = "") -> str:
    """Ponto único e robusto de extração de texto para currículos (PDF, DOCX, TXT, RTF, MD)."""
    ext = os.path.splitext(path.lower())[1]

    if ext == ".pdf" or "pdf" in content_type.lower():
        return extract_text_from_pdf(path)
    elif ext == ".docx" or "wordprocessingml" in content_type.lower():
        return extract_text_from_docx(path)
    elif ext in (".rtf",) or "rtf" in content_type.lower():
        return extract_text_from_rtf(path)
    elif ext in (".txt", ".md", ".text") or "plain" in content_type.lower() or "markdown" in content_type.lower():
        return extract_text_from_txt(path)
    else:
        # Detecta tipo pelo cabeçalho binário (magic bytes)
        try:
            with open(path, "rb") as f:
                head = f.read(16)
            if head.startswith(b"%PDF"):
                return extract_text_from_pdf(path)
            elif head.startswith(b"PK\x03\x04"):
                return extract_text_from_docx(path)
            elif head.startswith(b"{\\rtf"):
                return extract_text_from_rtf(path)
            else:
                return extract_text_from_txt(path)
        except Exception as e:
            if isinstance(e, ValueError):
                raise
            raise ValueError(f"Erro ao processar arquivo: {e}")


def _is_english_document(text: str) -> bool:
    """Identifica se o documento está predominantemente em inglês."""
    en_words = {"experience", "education", "skills", "summary", "university", "projects", "languages", "work", "years"}
    pt_words = {"experiência", "experiencia", "formação", "formacao", "habilidades", "resumo", "universidade", "projetos", "idiomas", "anos"}
    tokens = set(re.findall(r"\b\w+\b", text.lower()))
    en_hits = len(tokens.intersection(en_words))
    pt_hits = len(tokens.intersection(pt_words))
    return en_hits > pt_hits


def segment_resume(text: str) -> dict[str, str]:
    """Segmenta o currículo em seções semânticas baseadas em cabeçalhos comuns."""
    lines = text.splitlines()
    sections: dict[str, list[str]] = {
        "header": [],
        "summary": [],
        "experience": [],
        "education": [],
        "skills": [],
        "languages": [],
        "projects": [],
        "other": []
    }

    current_section = "header"
    header_found = False

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # Linha candidata a título de seção (curta, < 50 caracteres)
        matched_section = None
        if len(stripped) < 50:
            clean_line = stripped.lower().strip(":•-# ")
            for sec_name, patterns in SECTION_PATTERNS.items():
                for pat in patterns:
                    if re.fullmatch(pat, clean_line) or re.match(r"^" + pat + r"[:\s]*$", clean_line):
                        matched_section = sec_name
                        break
                if matched_section:
                    break

        if matched_section:
            current_section = matched_section
            header_found = True
            continue

        if not header_found and len(sections["header"]) > 5:
            current_section = "summary"

        sections[current_section].append(stripped)

    return {k: "\n".join(v) for k, v in sections.items()}


def _extract_summary(text: str, sections: dict[str, str]) -> str:
    """Extrai um resumo profissional conciso, sem dados pessoais de contato."""
    raw_summary = sections.get("summary") or ""

    if not raw_summary or len(raw_summary.strip()) < 30:
        for block in text.split("\n\n"):
            clean = block.strip()
            if len(clean) >= 60 and "@" not in clean and not re.search(r"\b(telefone|linkedin|github|celular|whatsapp)\b", clean, re.I):
                raw_summary = clean
                break

    lines = raw_summary.splitlines()
    filtered = []
    for line in lines:
        l = line.strip()
        if "@" in l or re.search(r"\b(https?://|linkedin\.com|github\.com|\+?\d{2,}\s*\d{4,})\b", l, re.I):
            continue
        if len(l) < 3:
            continue
        filtered.append(l)

    cleaned = " ".join(filtered).strip()
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned[:750]


def _extract_skills(text: str, sections: dict[str, str]) -> list[str]:
    """Extrai tecnologias do catálogo canônico e analisa listas e tags de seções de skills."""
    skills_found: dict[str, str] = {}
    lower_text = text.lower()

    # 1. Correspondência no catálogo com regex de limites de palavras
    for term, canonical in SKILLS_CANONICAL.items():
        if term in ("c++", "c#", ".net", "ci/cd", "12-factor", "12-factor app"):
            pattern = r"(?:^|[\s,;|\(/])" + re.escape(term) + r"(?:[\s,;|\)/]|$)"
        else:
            pattern = r"\b" + re.escape(term) + r"\b"

        if re.search(pattern, lower_text):
            skills_found[canonical.lower()] = canonical

    # 2. Extração dinâmica da seção de skills (lida com categorias tipo "Backend: Python, Django")
    skills_section = sections.get("skills", "")
    if skills_section:
        # Divide por linhas e depois por separadores comuns
        for raw_line in skills_section.splitlines():
            line = raw_line.strip()
            # Se a linha tem prefixo de categoria (ex: "Backend e Web: Django, FastAPI"), remove o prefixo
            if ":" in line and len(line.split(":", 1)[0]) < 35:
                line = line.split(":", 1)[1]

            items = re.split(r"[,;•|/\n]", line)
            for item in items:
                # Divide termos unidos por ' e ' ou ' and ' (ex: "Docker e Figma")
                subparts = re.split(r"\s+(?:e|and|&)\s+", item)
                for sp in subparts:
                    clean = sp.strip().strip("•-*[]()., ")
                    if 2 <= len(clean) <= 30 and not re.search(r"\b(experi[êe]ncia|anos|projetos?|conhecimento|trabalho|habilidade)\b", clean, re.I):
                        low_c = clean.lower()
                        if low_c in SKILLS_CANONICAL:
                            skills_found[SKILLS_CANONICAL[low_c].lower()] = SKILLS_CANONICAL[low_c]
                        elif clean.isupper() or any(c.isupper() for c in clean):
                            skills_found[low_c] = clean.title() if clean.islower() else clean

    # Ordena com tecnologias prioritárias no topo
    priority = ["Python", "Django", "Django ORM", "FastAPI", "SQL", "PostgreSQL", "SQLite", "APIs REST", "Docker", "Git", "Linux"]
    sorted_skills = []
    for p in priority:
        if p.lower() in skills_found:
            sorted_skills.append(skills_found.pop(p.lower()))
    sorted_skills.extend(sorted(skills_found.values()))
    return sorted_skills


def _extract_languages(text: str, sections: dict[str, str]) -> list[str]:
    """Extrai idiomas com nível de proficiência detectado."""
    target_text = (sections.get("languages") or "") + "\n" + text
    lower = target_text.lower()
    languages = []
    is_en_doc = _is_english_document(text)

    # Português
    if re.search(r"\b(portugu[êe]s|portuguese)\b", lower):
        if re.search(r"portugu[êe]s[^\n\.,;]*\b(nativo|nativa|materno|materna|fluente)\b", lower):
            languages.append("Português (Nativo)")
        else:
            languages.append("Português")
    elif not is_en_doc:
        languages.append("Português (Nativo)")

    # Inglês
    if re.search(r"\b(ingl[êe]s|english)\b", lower):
        if re.search(r"(ingl[êe]s|english)[^\n\.,;]*\b(fluente|fluent|nativo|native)\b", lower):
            languages.append("Inglês (Fluente)")
        elif re.search(r"(ingl[êe]s|english)[^\n\.,;]*\b(avan[çc]ado|advanced|c1|c2)\b", lower):
            languages.append("Inglês (Avançado)")
        elif re.search(r"(ingl[êe]s|english)[^\n\.,;]*\b(intermedi[áa]rio|intermediate|b1|b2|leitura\s+t[ée]cnica)\b", lower):
            languages.append("Inglês (Intermediário - Leitura técnica)")
        elif re.search(r"(ingl[êe]s|english)[^\n\.,;]*\b(b[áa]sico|basic|a1|a2|iniciante)\b", lower):
            languages.append("Inglês (Básico)")
        else:
            languages.append("Inglês")
    elif is_en_doc:
        languages.append("Inglês (Fluente)")

    # Espanhol
    if re.search(r"\b(espanhol|spanish)\b", lower):
        if re.search(r"(espanhol|spanish)[^\n\.,;]*\b(fluente|fluent|nativo)\b", lower):
            languages.append("Espanhol (Fluente)")
        elif re.search(r"(espanhol|spanish)[^\n\.,;]*\b(avan[çc]ado|advanced)\b", lower):
            languages.append("Espanhol (Avançado)")
        elif re.search(r"(espanhol|spanish)[^\n\.,;]*\b(intermedi[áa]rio|intermediate)\b", lower):
            languages.append("Espanhol (Intermediário)")
        else:
            languages.append("Espanhol (Básico)")

    return languages


def _extract_experience_and_seniority(text: str, sections: dict[str, str]) -> tuple[int, str]:
    """Calcula anos de experiência e senioridade sem contaminação por níveis de idioma."""
    non_language_text = f"{sections.get('header', '')}\n{sections.get('summary', '')}\n{sections.get('experience', '')}\n{sections.get('education', '')}\n{sections.get('projects', '')}".lower()

    # Indicadores claros de estudante / estágio / transição
    is_student_or_intern = bool(re.search(
        r"\b(cursando|graduando|estudante\s+de|previs[ãa]o\s+de\s+formatura|bacharelado\s+em\s+andamento|est[áa]gio|estagi[áa]ri[oa]|trainee)\b",
        non_language_text
    ))

    # Anos de experiência explícitos (ex: "3 anos de experiência")
    exp_explicit = re.findall(r"(\d+)\s*(?:\+|anos?|years?|yrs?)\s*(?:de\s+)?(?:experi[êe]ncia|experience)", non_language_text)
    years = 0
    if exp_explicit:
        years = max(int(y) for y in exp_explicit)
    else:
        # Extração por intervalos de datas em experiências profissionais
        years_found = [int(y) for y in re.findall(r"\b(20[0-2]\d)\b", sections.get("experience", "") or non_language_text)]
        if len(years_found) >= 2:
            min_y, max_y = min(years_found), max(years_found)
            now_y = datetime.now().year
            diff = min(now_y - min_y, max_y - min_y)
            if 0 <= diff <= 30:
                years = max(years, diff)

    # Classificação de senioridade
    if re.search(r"\b(s[eê]nior|sr\.?|lead|tech\s+lead|especialista|principal)\b", non_language_text):
        seniority = "senior"
        years = max(years, 5)
    elif re.search(r"\b(pleno|mid-level|pl\.?)\b", non_language_text):
        seniority = "mid"
        years = max(years, 3)
    elif is_student_or_intern or years <= 2 or re.search(r"\b(j[úu]nior|jr\.?|iniciante|entry-level)\b", non_language_text):
        seniority = "junior"
        years = min(years, 2)
        if years == 0 and is_student_or_intern:
            years = 1
    else:
        seniority = "mid" if years >= 3 else "junior"

    return years, seniority


def _extract_headline_and_roles(text: str, sections: dict[str, str], skills: list[str], seniority: str) -> tuple[str, list[str]]:
    """Extrai headline profissional e gera lista de cargos desejados compatíveis com múltiplos perfis."""
    header_text = sections.get("header", "")
    lines = [l.strip() for l in header_text.splitlines() if l.strip()]

    found_headline = ""
    # Candidatos a headline no topo
    for line in lines[1:5]:
        if not any(c in line for c in "@•/|") and len(line) < 70:
            if re.search(r"\b(desenvolvedor[a]?|developer|engenheir[oa]|software|analista|backend|frontend|full\s*stack|estagi[áa]ri[oa]|engineer|architect|devops|cloud|data|sre)\b", line, re.I):
                found_headline = line
                break

    if not found_headline:
        m = re.search(r"(?:objetivo|cargo\s+desejado|área\s+de\s+interesse)[:\-]?\s*([^\n\.]+)", sections.get("summary", "") or text, re.I)
        if m:
            cand = m.group(1).strip()
            if 5 < len(cand) < 65:
                found_headline = cand

    has_python = "Python" in skills
    has_django = "Django" in skills
    has_react = "React" in skills
    has_node = "Node.js" in skills
    has_java = "Java" in skills
    has_devops = any(s in skills for s in ("Docker", "Kubernetes", "AWS", "Terraform", "CI/CD"))

    sen_suffix = "Júnior" if seniority == "junior" else ("Pleno" if seniority == "mid" else "Sênior")

    if not found_headline:
        if has_python and has_django:
            found_headline = "Desenvolvedor Backend Python | Django | APIs REST"
        elif has_python:
            found_headline = f"Desenvolvedor Python {sen_suffix} | Backend"
        elif has_react and has_node:
            found_headline = f"Desenvolvedor Full Stack {sen_suffix} | React & Node.js"
        elif has_react:
            found_headline = f"Desenvolvedor Frontend {sen_suffix} | React"
        elif has_java:
            found_headline = f"Desenvolvedor Java {sen_suffix} | Spring Boot"
        elif has_devops:
            found_headline = f"Engenheiro DevOps {sen_suffix} | Cloud"
        else:
            found_headline = f"Desenvolvedor de Software {sen_suffix}"

    # Gera lista de cargos compatíveis com a senioridade e especialidade
    roles = []
    low_head = found_headline.lower()

    if "python" in low_head or has_python:
        if seniority == "junior":
            roles.extend([
                "Desenvolvedor Python Júnior",
                "Desenvolvedor Backend Júnior",
                "Desenvolvedor de Software Júnior",
                "Estagiário de Desenvolvimento de Software",
                "Estágio em TI",
                "Junior Python Developer",
                "Desenvolvedor Backend"
            ])
        elif seniority == "mid":
            roles.extend([
                "Desenvolvedor Python Pleno",
                "Desenvolvedor Backend Pleno",
                "Desenvolvedor de Software Pleno",
                "Engenheiro de Software",
                "Backend Developer"
            ])
        else:
            roles.extend([
                "Desenvolvedor Python Sênior",
                "Desenvolvedor Backend Sênior",
                "Senior Python Developer",
                "Tech Lead Python"
            ])
    elif "front" in low_head or has_react:
        roles.extend([
            f"Desenvolvedor Frontend {sen_suffix}",
            f"Desenvolvedor React {sen_suffix}",
            "Desenvolvedor Web",
            "Engenheiro de Software"
        ])
    elif "devops" in low_head or "cloud" in low_head:
        roles.extend([
            f"Engenheiro DevOps {sen_suffix}",
            f"Analista Cloud {sen_suffix}",
            "SRE",
            "Engenheiro de Infraestrutura"
        ])
    else:
        roles.extend([
            found_headline,
            f"Desenvolvedor de Software {sen_suffix}",
            f"Desenvolvedor Backend {sen_suffix}",
            "Engenheiro de Software"
        ])

    return found_headline, list(dict.fromkeys(roles))[:8]


def deterministic_parse_resume(text: str) -> dict:
    """
    Parser determinístico universal e inteligente para currículos.
    Compatível com formatos de coluna única, coluna dupla, funcionais, acadêmicos e internacionais.
    """
    if not text:
        return {}

    sections = segment_resume(text)
    skills = _extract_skills(text, sections)
    years, seniority = _extract_experience_and_seniority(text, sections)
    languages = _extract_languages(text, sections)
    headline, roles = _extract_headline_and_roles(text, sections, skills, seniority)
    summary = _extract_summary(text, sections)

    return {
        "headline": headline,
        "summary": summary,
        "years_experience": years,
        "seniority": seniority,
        "skills": skills,
        "roles": roles,
        "languages": languages
    }
